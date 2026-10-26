"""
File parser tool — turns an uploaded lecture file into plain text.

Supports: .pdf, .pptx, .txt, .md. Used by the Concept Extractor agent.
Enforces a character cap so a huge upload can't blow up token spend
(part of the overspend defense).
"""
from __future__ import annotations

from pathlib import Path

# Overspend guard: never feed more than this many characters to the model.
MAX_CHARS = 60_000

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".txt", ".md"}


class UnsupportedFileType(ValueError):
    pass


def _parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _parse_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def extract_text(path: str | Path, *, max_chars: int = MAX_CHARS) -> str:
    """Extract text from a lecture file, truncated to max_chars."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = _parse_pdf(path)
    elif ext == ".pptx":
        text = _parse_pptx(path)
    elif ext in (".txt", ".md"):
        text = _parse_text(path)
    else:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[... truncated to protect token budget ...]"
    return text
